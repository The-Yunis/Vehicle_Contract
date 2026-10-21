"""ARACLAR.xlsx'teki 'Üret = E' satırları için SABLON.docx'ten sözleşme üretir.

Kullanım:  python uret.py [excel_yolu] [cikti_klasoru]
"""
import re
import sys
import time
from datetime import date, datetime
from pathlib import Path

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import load_workbook

from ortak import ALAN, metin_degistir, paragraf_sil

KOK = Path(__file__).resolve().parent.parent
EXCEL = Path(sys.argv[1]) if len(sys.argv) > 1 else KOK / "ARACLAR.xlsx"
CIKTI = Path(sys.argv[2]) if len(sys.argv) > 2 else KOK / "CIKTI"
SABLON = KOK / "SABLON.docx"
BOS_TARIH = "..../..../20...."
KUTU = "☐"

# ---------------------------------------------------------------- yardımcılar

BIR = ["", "bir", "iki", "üç", "dört", "beş", "altı", "yedi", "sekiz", "dokuz"]
ON = ["", "on", "yirmi", "otuz", "kırk", "elli", "altmış", "yetmiş", "seksen", "doksan"]
BOLUK = ["", "bin", "milyon", "milyar"]


def yaziyla(n):
    if n == 0:
        return "sıfır"
    parcalar = []
    i = 0
    while n:
        n, uc = divmod(n, 1000)
        if uc:
            y, kalan = divmod(uc, 100)
            s = ("" if y == 0 else "yüz" if y == 1 else BIR[y] + "yüz") + ON[kalan // 10] + BIR[kalan % 10]
            if i == 1 and uc == 1:
                s = ""
            parcalar.append(s + BOLUK[i])
        i += 1
    return "".join(reversed(parcalar))


def tl_yazi(tutar):
    kurus_top = round(tutar * 100)
    lira, kurus = divmod(kurus_top, 100)
    s = f"{yaziyla(lira)} Türk Lirası"
    return s + (f" {yaziyla(kurus)} kuruş" if kurus else "")


def tl(tutar):
    s = f"{tutar:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{s} TL"


def metin(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def tarih(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = metin(v)
    for bicim in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, bicim).date()
        except ValueError:
            pass
    return None


def tablo_oku(ws):
    satirlar = list(ws.iter_rows(values_only=True))
    basliklar = [metin(b) for b in satirlar[0]]
    for no, satir in enumerate(satirlar[1:], 2):
        kayit = dict(zip(basliklar, satir))
        if any(metin(v) for v in satir):
            yield no, kayit


def sutun(kayit, bas):
    """Başlığı `bas` ile başlayan sütunun değeri (başlık metni sonradan biraz değişse de bulur)."""
    for k, v in kayit.items():
        if k.startswith(bas):
            return v
    return None


# ---------------------------------------------------------------- word işleri

def kenarlik(tablo):
    tblPr = tablo._tbl.tblPr
    kenar = OxmlElement("w:tblBorders")
    for ad in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{ad}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "000000")
        kenar.append(e)
    tblPr.append(kenar)


def hucre(c, yazi, kalin=False, boyut=8.5, orta=False, golge=False):
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if orta else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(yazi)
    r.bold = kalin or None
    r.font.size = Pt(boyut)
    if golge:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "D9D9D9")
        c._tc.get_or_add_tcPr().append(shd)


def tablo(doc, satirlar, genislikler, baslik=True, kalin_ilk_sutun=False, boyut=8.5):
    t = doc.add_table(rows=len(satirlar), cols=len(genislikler))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    kenarlik(t)
    for i, satir in enumerate(satirlar):
        for j, deger in enumerate(satir):
            c = t.cell(i, j)
            c.width = Cm(genislikler[j])
            ilk = baslik and i == 0
            hucre(c, deger, kalin=ilk or (kalin_ilk_sutun and j == 0), boyut=boyut,
                  orta=ilk, golge=ilk or (kalin_ilk_sutun and j == 0))
    return t


def arac_tablosu(doc, araclar, toplam):
    satirlar = [["#", "Plaka", "Cinsi", "Marka / Model", "Yıl", "Şasi No",
                 "Teknik Özellik / Ataşman", "Aylık Kira (KDV Hariç)"]]
    for i, a in enumerate(araclar, 1):
        ozellik = "\n".join(x for x in (a["ozellik"], a["atasman"] and f'Ataşman: {a["atasman"]}') if x)
        satirlar.append([str(i), a["plaka"] or "-", a["cins"], f'{a["marka"]} {a["model"]}',
                         a["yil"] or "-", a["sasi"], ozellik or "-", tl(a["bedel"])])
    satirlar.append(["", "", "", "", "", "", "TOPLAM", tl(toplam)])
    t = tablo(doc, satirlar, [0.6, 1.7, 1.9, 2.1, 1.0, 3.4, 2.9, 2.4])
    for c in t.rows[-1].cells[-2:]:
        for r in c.paragraphs[0].runs:
            r.bold = True
    return t


def tutanaklar(doc, d, araclar):
    for i, a in enumerate(araclar, 1):
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        r = p.add_run(f"EK-1/{i}: İŞ MAKİNASI TESLİM VE İADE TUTANAĞI")
        r.bold, r.font.size = True, Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(
            f'Sözleşme No: {d["SOZ_NO"]}    Kiralayan: {d["K_AD"]}    Kiracı: {d["R_UNVAN"]}'
        ).runs[0].font.size = Pt(9)

        tablo(doc, [
            ["Plaka", a["plaka"] or "-", "Şasi No", a["sasi"]],
            ["Cinsi", a["cins"], "Marka / Model", f'{a["marka"]} {a["model"]}'],
            ["Model Yılı", a["yil"] or "-", "Aylık Kira", tl(a["bedel"]) + " + KDV"],
            ["Teknik Özellik", a["ozellik"] or "-", "Ataşmanlar", a["atasman"] or "-"],
        ], [2.5, 5.5, 2.5, 5.5], baslik=False, kalin_ilk_sutun=False, boyut=9)
        doc.add_paragraph()

        atasmanlar = [x.strip() for x in a["atasman"].split(",") if x.strip()] or ["................"]
        ata = "\n".join(f"{KUTU} {x}" for x in atasmanlar)
        belgeler = "\n".join(f"{KUTU} {x}" for x in (
            "Ruhsat / Tescil Belgesi", "Zorunlu Mali Sorumluluk (Trafik) Poliçesi", "Kasko Poliçesi",
            "Makine Kırılması Poliçesi", "İMM Poliçesi (kiralama şerhli)", "Muayene / Periyodik Kontrol",
            "Kullanım Talimatnamesi"))
        bos = "\n\n"
        tablo(doc, [
            ["KONTROL", "TESLİM", "İADE"],
            ["Tarih / Saat", "...../...../20.....   ....:....", "...../...../20.....   ....:...."],
            ["Teslim / İade Yeri", d["TESLIM_YERI"], ""],
            ["Çalışma Saati Sayacı", "", ""],
            ["Kilometre", "", ""],
            ["Yakıt Seviyesi", "", ""],
            ["Lastik Durumu", "", ""],
            ["Mevcut Hasar / Eksikler", bos, bos],
            ["Ataşmanlar", ata, ata],
            ["Teslim Edilen Belgeler", belgeler, belgeler],
            ["Fotoğraf", "........ adet fotoğraf çekildi", "........ adet fotoğraf çekildi"],
            ["Makina çalışır ve operasyonel", f"{KUTU} Evet   {KUTU} Hayır", f"{KUTU} Evet   {KUTU} Hayır"],
            ["Açıklama", bos, bos],
        ], [4.0, 6.0, 6.0], kalin_ilk_sutun=True)

        not_p = doc.add_paragraph()
        not_p.add_run(
            "Not: Kira bedeli, makinanın bu tutanakta çalışır ve operasyonel olarak teslim alındığı tarihten "
            "itibaren işler (Sözleşme Madde 2). Tuz/kimyasal kaynaklı korozyon ile bıçak ve lastik aşınması "
            "olağan aşınmadır (Madde 6.1.5).").font.size = Pt(8)
        not_p.runs[0].italic = True

        tablo(doc, [
            ["TESLİM EDEN (Kiralayan)", "TESLİM ALAN (Kiracı)", "İADE EDEN (Kiracı)", "İADE ALAN (Kiralayan)"],
            ["Ad Soyad:\n\nİmza:\n\n", "Ad Soyad:\n\nİmza:\n\n", "Ad Soyad:\n\nİmza:\n\n", "Ad Soyad:\n\nİmza:\n\n"],
        ], [4.0, 4.0, 4.0, 4.0], boyut=8)


def tum_paragraflar(doc):
    def tablolar(tablolar_):
        for t in tablolar_:
            for satir in t.rows:
                for c in satir.cells:
                    yield from c.paragraphs
    yield from doc.paragraphs
    yield from tablolar(doc.tables)
    for bolum in doc.sections:  # altbilgide {{SOZ_NO}} var
        for kisim in (bolum.header, bolum.footer):
            if not kisim.is_linked_to_previous:
                yield from kisim.paragraphs
                yield from tablolar(kisim.tables)


LOGO_ALANI = 1.8   # cm²: her logonun sayfa başlığında kaplayacağı yaklaşık alan


def logo_orani(yol):
    """PNG başlığından en/boy oranı (ek kütüphane gerektirmeden)."""
    import struct
    with open(yol, "rb") as f:
        f.seek(16)
        en, boy = struct.unpack(">II", f.read(8))
    return en / boy


def logo_basligi(doc, logolar):
    """Sayfa başlığına logoları yan yana (sol/sağ) yerleştirir."""
    bolum = doc.sections[0]
    bas = bolum.header
    bas.is_linked_to_previous = False
    genislik = bolum.page_width - bolum.left_margin - bolum.right_margin
    t = bas.add_table(rows=1, cols=len(logolar), width=genislik)
    for i, (c, logo) in enumerate(zip(t.rows[0].cells, logolar)):
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        # Logoların en-boy oranı farklı (Öztaç 2,85 : 1, HİKA 4,32 : 1); aynı yükseklikte HİKA çok
        # daha iri görünür. Kapladıkları alanı eşitlemek için yükseklik oranın kareköküyle ölçeklenir.
        p.add_run().add_picture(str(logo), height=Cm((LOGO_ALANI / logo_orani(logo)) ** 0.5))
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    ilk = bas.paragraphs[0]
    if not ilk.text:
        ilk._p.getparent().remove(ilk._p)


def sozlesme_yaz(d, araclar, toplam, hedef, sablon=None, logolar=()):
    doc = docx.Document(sablon or SABLON)
    if logolar:
        logo_basligi(doc, logolar)

    for p in list(tum_paragraflar(doc)):
        for alan in set(ALAN.findall(p.text)):
            ad = alan.strip("{}?")
            if alan.startswith("{{?") and not d.get(ad):
                paragraf_sil(p)
                break
            if ad in ("ARAC_TABLOSU", "TUTANAKLAR"):
                continue
            if ad not in d:
                raise KeyError(f"Şablonda bilinmeyen alan: {alan}")
            metin_degistir(p, alan, d[ad])

    for p in doc.paragraphs:
        if "{{ARAC_TABLOSU}}" in p.text:
            t = arac_tablosu(doc, araclar, toplam)
            p._p.addnext(t._tbl)
            paragraf_sil(p)
            break
    for p in doc.paragraphs:
        if "{{TUTANAKLAR}}" in p.text:
            paragraf_sil(p)
            tutanaklar(doc, d, araclar)
            break

    kalan = [p.text for p in tum_paragraflar(doc) if ALAN.search(p.text)]
    if kalan:
        raise ValueError(f"Doldurulmamış alan kaldı: {kalan[0][:80]}")
    doc.save(hedef)


def tckn_gecerli(s):
    if not re.fullmatch(r"[1-9]\d{10}", s):
        return False
    d = [int(c) for c in s]
    return (sum(d[0:9:2]) * 7 - sum(d[1:8:2])) % 10 == d[9] and sum(d[:10]) % 10 == d[10]


def kiraci_bilgisi(ayar):
    """Ayarlar (başlık -> değer) sözlüğünden Kiracı alanlarını ve genel uyarıları çıkarır."""
    R = {
        "R_UNVAN": metin(ayar.get("Kiracı Unvanı")),
        "R_ADRES": metin(ayar.get("Kiracı Adresi")),
        "R_TEL": metin(ayar.get("Kiracı Telefon")),
        "R_EPOSTA": metin(ayar.get("Kiracı E-posta")),
        "R_VD": metin(ayar.get("Kiracı Vergi Dairesi")),
        "R_VN": metin(ayar.get("Kiracı VKN")),
        "R_YETKILI": metin(sutun(ayar, "Kiracı İmza Yetkilisi")),
        "SEVK_SURESI": metin(sutun(ayar, "Sevk Talebinden")) or "48 (kırk sekiz) saat",
    }
    uyarilar = [f"Ayarlar: Kiracı {ad} boş." for k, ad in (
        ("R_TEL", "telefon"), ("R_EPOSTA", "e-posta"), ("R_YETKILI", "imza yetkilisi")) if R[k] in ("", "-")]
    return R, uyarilar


def sozlesme_isle(s, arac_kayitlari, R, kdv_orani, uyarilar_genel=(), cikti=None, sablon=None, logolar=()):
    """Tek sözleşmeyi doğrular ve üretir (PDF hariç).

    s: sözleşme alanları (Excel başlıkları -> değer); arac_kayitlari: [(etiket, araç alanları)].
    Dönüş: {"ok": bool, "soz_no", "hatalar", "uyarilar", "hedef", "adet", "genel"}
    """
    cikti = Path(cikti or CIKTI)
    soz_no = metin(sutun(s, "Sözleşme No"))
    hatalar, uyarilar = [], list(uyarilar_genel)
    sonuc = {"ok": False, "soz_no": soz_no, "hatalar": hatalar, "uyarilar": uyarilar}

    if not R["R_UNVAN"] or not R["R_ADRES"]:
        hatalar.append("Ayarlar: Kiracı unvanı ve adresi zorunlu")
    k_ad = metin(sutun(s, "Kiralayan Adı"))
    kimlik = re.sub(r"\D", "", metin(sutun(s, "TCKN")))
    baslangic, bitis = tarih(sutun(s, "Teslim (Başlangıç)")), tarih(sutun(s, "Bitiş"))
    imza = tarih(sutun(s, "İmza Tarihi"))
    teslim_yeri = metin(sutun(s, "Teslim Yeri"))
    if not soz_no:
        hatalar.append("Sözleşme No boş")
    if not k_ad:
        hatalar.append("Kiralayan adı boş")
    if len(kimlik) == 11 and not tckn_gecerli(kimlik):
        hatalar.append("TC Kimlik No hatalı (kontrol hanesi tutmuyor)")
    elif len(kimlik) not in (10, 11):
        hatalar.append("TCKN 11, VKN 10 haneli olmalı")
    if not metin(sutun(s, "Kiralayan Adresi")):
        hatalar.append("Kiralayan adresi boş")
    if not baslangic:
        hatalar.append("Teslim (başlangıç) tarihi boş/geçersiz")
    if not bitis:
        hatalar.append("Bitiş tarihi boş/geçersiz")
    if baslangic and bitis and bitis <= baslangic:
        hatalar.append("Bitiş tarihi başlangıçtan sonra olmalı")
    if not teslim_yeri:
        hatalar.append("Teslim yeri boş")
    banka = metin(sutun(s, "Banka"))
    iban = re.sub(r"\s", "", metin(sutun(s, "IBAN"))).upper()
    # Ödeme: vadeli çek (şirketin olağan yöntemi, vade kiralayanla ayrı anlaşılır) ya da havale/EFT.
    # Eski kayıtlarda ödeme şekli yoksa havale sayılır (o zamanki şablon IBAN'a ödeme diyordu).
    odeme = metin(sutun(s, "Ödeme Şekli")).lower() or "havale"
    vade = re.sub(r"\D", "", metin(sutun(s, "Çek Vadesi")))
    if odeme == "cek":
        if not vade or not 1 <= int(vade) <= 365:
            hatalar.append("Çek vadesi 1 ile 365 gün arasında olmalı")
        if iban and not re.fullmatch(r"TR\d{24}", iban):
            hatalar.append("IBAN geçersiz (TR + 24 rakam olmalı) — çekle ödemede boş bırakılabilir")
    else:
        if not banka:
            hatalar.append("Banka adı boş")
        if not re.fullmatch(r"TR\d{24}", iban):
            hatalar.append("IBAN geçersiz (TR + 24 rakam olmalı)")
    iban = " ".join(iban[i:i + 4] for i in range(0, len(iban), 4))
    if metin(sutun(s, "E-posta")) in ("", "-"):
        uyarilar.append("Kiralayan e-postası boş (Madde 1.4 bildirimleri e-posta ile yapılıyor)")
    if metin(sutun(s, "Telefon")) in ("", "-"):
        uyarilar.append("Kiralayan telefonu boş")

    liste = []
    for etiket, a in arac_kayitlari:
        arac = {
            "plaka": metin(sutun(a, "Plaka")).upper(), "cins": metin(sutun(a, "Cinsi")),
            "marka": metin(sutun(a, "Marka")), "model": metin(sutun(a, "Model")),
            "yil": metin(sutun(a, "Model Yılı")), "sasi": metin(sutun(a, "Şasi")).upper(),
            "atasman": metin(sutun(a, "Ataşman")), "ozellik": metin(sutun(a, "Teknik Özellik")),
        }
        try:
            arac["bedel"] = float(sutun(a, "Aylık Kira") or 0)
        except (TypeError, ValueError):
            arac["bedel"] = 0
        eksik = [ad for ad, k in (("Cinsi", "cins"), ("Marka", "marka"), ("Model", "model"),
                                   ("Şasi No", "sasi")) if not arac[k]]
        if arac["bedel"] <= 0:
            eksik.append("Aylık Kira")
        if eksik:
            hatalar.append(f"{etiket}: {', '.join(eksik)} eksik")
        if not arac["plaka"]:
            uyarilar.append(f"{etiket}: plaka boş (tescilsiz makina ise sorun yok)")
        liste.append(arac)
    if not liste:
        hatalar.append("Bu sözleşmeye bağlı araç yok")
    if hatalar:
        return sonuc

    toplam = sum(a["bedel"] for a in liste)
    kdv = round(toplam * kdv_orani / 100, 2)
    d = dict(R)
    d.update({
        "SOZ_NO": soz_no, "K_AD": k_ad,
        "K_KIMLIK": f"T.C. Kimlik No: {kimlik}" if len(kimlik) == 11 else f"VKN: {kimlik}",
        "K_ADRES": metin(sutun(s, "Kiralayan Adresi")), "K_TEL": metin(sutun(s, "Telefon")),
        "K_EPOSTA": metin(sutun(s, "E-posta")), "K_VD": metin(sutun(s, "Vergi Dairesi")),
        "K_VN": kimlik if len(kimlik) == 10 else "",
        "K_YETKILI": metin(sutun(s, "Şirket Yetkilisi")), "K_BANKA": banka, "K_IBAN": iban,
        "BASLANGIC": baslangic.strftime("%d.%m.%Y"), "BITIS": bitis.strftime("%d.%m.%Y"),
        "TESLIM_YERI": teslim_yeri,
        "IMZA_TARIHI": imza.strftime("%d.%m.%Y") if imza else BOS_TARIH,
        "ARA_TOPLAM": tl(toplam), "ARA_TOPLAM_YAZI": tl_yazi(toplam),
        "KDV_ORANI": metin(kdv_orani), "KDV": tl(kdv),
        "GENEL_TOPLAM": tl(toplam + kdv), "GENEL_TOPLAM_YAZI": tl_yazi(toplam + kdv),
    })
    d["K_IMZACI"] = d["K_YETKILI"] or k_ad
    d["R_IMZACI"] = d["R_YETKILI"] or "Ad Soyad:"
    if odeme == "cek":
        gun = int(vade)
        d["ODEME_SEKLI"] = f"Kiralayan adına düzenlenecek {gun} ({yaziyla(gun)}) gün vadeli çek ile"
        d["ODEME_EK"] = (
            "Çekin Kiralayana veya yetkilisine teslimiyle Kiracının ilgili döneme ilişkin ödeme yükümlülüğü, "
            "çek bedelinin vadesinde ödenmesi kaydıyla yerine getirilmiş sayılır. Kiralayan çek vadesinden önce "
            "ödeme talep edemez; vade farkı, faiz veya sair herhangi bir bedel isteyemez."
            + (f" Kiralayanın banka hesabı: {banka} – {iban}." if iban else ""))
    else:
        d["ODEME_SEKLI"] = "banka havalesi/EFT yoluyla"
        d["ODEME_EK"] = (
            f"Kira bedelleri yalnızca Kiralayan adına açılmış {banka} nezdindeki {iban} IBAN numaralı hesaba "
            "ödenir. Kiralayan, hesap bilgisindeki değişikliği imzalı yazı ile bildirmedikçe bu hesaba yapılan "
            "ödemeler geçerli ve borçtan kurtarıcıdır; üçüncü kişi adına açılmış hesaplara ödeme talep edilemez.")

    cikti.mkdir(exist_ok=True)
    guvenli = re.sub(r'[\\/:*?"<>|]+', "_", f"{soz_no}_{k_ad}")[:80].rstrip(". ")
    hedef = cikti / f"{guvenli}.docx"
    try:
        sozlesme_yaz(d, liste, toplam, hedef, sablon, logolar)
    except PermissionError:
        hatalar.append(f"{hedef.name} Word'de açık, kapatıp tekrar deneyin")
        return sonuc
    sonuc.update(ok=True, hedef=hedef, adet=len(liste), genel=toplam + kdv, toplam=toplam, kdv=kdv,
                 araclar=liste, alanlar=d)
    return sonuc


# ---------------------------------------------------------------- ana akış (Excel)

def main():
    if not EXCEL.exists():
        sys.exit(f"Excel bulunamadı: {EXCEL}")
    wb = load_workbook(EXCEL, data_only=True)
    ayar = {metin(r[0]): r[1] for r in wb["Ayarlar"].iter_rows(values_only=True) if r[0]}
    kdv_orani = float(ayar.get("KDV Oranı (%)") or 20)
    R, uyarilar_genel = kiraci_bilgisi(ayar)

    araclar = {}
    for no, a in tablo_oku(wb["Araçlar"]):
        araclar.setdefault(metin(sutun(a, "Sözleşme No")), []).append((f"Araçlar satır {no}", a))

    sonuclar = []
    for no, s in tablo_oku(wb["Sözleşmeler"]):
        if metin(sutun(s, "Üret")).upper() == "E":
            sonuclar.append(sozlesme_isle(s, araclar.get(metin(sutun(s, "Sözleşme No")), []),
                                          R, kdv_orani, uyarilar_genel))

    pdf_hatasi = pdf_yap([s["hedef"] for s in sonuclar if s["ok"]])
    if pdf_hatasi:
        print(pdf_hatasi)

    print("=" * 70)
    for s in sonuclar:
        if s["ok"]:
            print(f'✔ {s["soz_no"]}: {s["adet"]} araç, KDV dahil {tl(s["genel"])}  ->  CIKTI\\{s["hedef"].name}')
            for u in dict.fromkeys(s["uyarilar"]):
                print(f"     ! Uyarı: {u}")
        else:
            print(f'✘ {s["soz_no"]} ÜRETİLMEDİ:')
            for h in s["hatalar"]:
                print(f"     - {h}")
    if not sonuclar:
        print("'Üret (E/H)' sütununda E olan sözleşme yok.")
    print("=" * 70)


PDF_IPUCU = ("PDF oluşturulamadı; Word dosyası hazır. Word'de açık bir uyarı penceresi kalmış olabilir: "
             "Word'ü bir kez açıp kapatın, sonra sözleşmeyi yeniden oluşturun.")


def _word_ile_pdf(dosyalar):
    """Tek deneme: yeni bir Word örneğiyle PDF'leri üretir. Hata olursa istisna fırlatır."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()  # sunucu iş parçacığından çağrılabiliyor
    word = win32com.client.DispatchEx("Word.Application")
    try:
        word.Visible = False
        word.DisplayAlerts = 0  # wdAlertsNone: kaydetme/dönüştürme uyarıları otomasyonu kilitlemesin
        for f in dosyalar:
            belge = word.Documents.Open(str(Path(f).resolve()), ConfirmConversions=False, ReadOnly=True,
                                        AddToRecentFiles=False, Visible=False)
            try:
                belge.SaveAs2(str(Path(f).with_suffix(".pdf").resolve()), FileFormat=17)
            finally:
                belge.Close(False)
    finally:
        try:
            word.Quit(False)
        except Exception:
            pass  # Word yanıt vermiyorsa kapatma hatası asıl sonucu gölgelemesin


def pdf_yap(dosyalar, deneme=2, gunluk=None):
    """Word ile PDF üretir. Sorun olursa kullanıcıya gösterilecek metni döndürür (docx'ler yine de hazırdır).
    gunluk: asıl hatayı kaydetmek için çağrılacak fonksiyon (ör. kms.gunluge_yaz)."""
    if not dosyalar:
        return None
    son_hata = None
    for i in range(deneme):
        try:
            _word_ile_pdf(dosyalar)
            if all(Path(f).with_suffix(".pdf").exists() for f in dosyalar):
                return None
            son_hata = RuntimeError("Word PDF dosyasını oluşturmadı")
        except ImportError:
            return "PDF oluşturulamadı: bu bilgisayarda Microsoft Word bulunamadı. Word dosyası hazır."
        except Exception as e:
            son_hata = e
        time.sleep(2)  # ikinci deneme yeni bir Word örneğiyle
    if gunluk:
        gunluk(f"PDF oluşturulamadı ({deneme} deneme): {son_hata!r}")
    return PDF_IPUCU


if __name__ == "__main__":
    main()
