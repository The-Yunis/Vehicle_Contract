"""Şirketin verdiği orijinal sözleşmeden, Kiracı lehine düzeltilmiş ve
{{ALAN}} yer tutuculu SABLON.docx dosyasını üretir.

Bir kez çalıştırılır. Sonrasında metin değişiklikleri doğrudan SABLON.docx
üzerinde Word ile yapılabilir ({{...}} alanlarına dokunmadan).
"""
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from bicim import kurumsal_bicim
from ortak import metin_degistir, paragraf_ekle_sonra, paragraf_sil, paragraf_yaz

KOK = Path(__file__).resolve().parent.parent
ORIJINAL = next(KOK.glob("İŞ MAKİNASI KİRALAMA SÖZLEŞMESİ*.docx"))
HEDEF = KOK / "SABLON.docx"


def bul(doc, bas, sira=0):
    """Metni `bas` ile başlayan `sira`'ncı paragrafı döndürür."""
    bulunan = [p for p in doc.paragraphs if p.text.strip().startswith(bas)]
    if len(bulunan) <= sira:
        sys.exit(f"Paragraf bulunamadı: {bas!r}")
    return bulunan[sira]


def main():
    doc = docx.Document(ORIJINAL)

    # --- MADDE 1: TARAFLAR -------------------------------------------------
    paragraf_yaz(bul(doc, "KİRALAYAN "), [
        ("KİRALAYAN              :  ", True), ("{{K_AD}}", False), ("   {{K_KIMLIK}}", False)])
    adres_k = bul(doc, "Adresi", 0)
    paragraf_yaz(adres_k, [
        ("Adresi                    :  ", True), ("{{K_ADRES}} ", False),
        ('(Bundan böyle "KİRALAYAN" olarak anılacaktır.)', False)])
    for sira, onek in ((0, "K"), (1, "R")):
        for etiket, alan in (("Telefon", "TEL"), ("E-posta", "EPOSTA"),
                             ("Vergi Dairesi", "VD"), ("Vergi No", "VN")):
            bul(doc, etiket, sira).add_run(f"{{{{{onek}_{alan}}}}}")
    # Şirket kiralayanlarda imza yetkilisi; şahısta satır kendiliğinden silinir.
    paragraf_ekle_sonra(bul(doc, "Vergi No", 0), [
        ("Yetkili\t\t: ", True), ("{{?K_YETKILI}}", False)])

    paragraf_yaz(bul(doc, "KİRACI "), [
        ("KİRACI                    :  ", True), ("{{R_UNVAN}}", True)])
    paragraf_yaz(bul(doc, "Adresi", 1), [
        ("Adresi                    :  ", True), ("{{R_ADRES}} ", False),
        ('(Bundan böyle "KİRACI" olarak anılacaktır.)', False)])
    paragraf_ekle_sonra(bul(doc, "Vergi No", 1), [
        ("Yetkili\t\t: ", True), ("{{?R_YETKILI}}", False)])

    metin_degistir(bul(doc, "1.2."), "anılacakladır", "anılacaklardır")

    # --- MADDE 2: TANIMLAR -------------------------------------------------
    paragraf_yaz(bul(doc, "Kira Süresi:"), [
        ("Kira Süresi:", True),
        (" Kira Ödemelerinin Başlangıç Tarihi ile {{BITIS}} tarihi arasındaki dönemi ifade eder. "
         "Ancak Kiracı'nın aşağıda tanımlanan \"İşin Bitmesi\" nedeniyle erken fesih hakkı saklıdır.", False)])
    paragraf_yaz(bul(doc, "Teslim Adresi/yeri:"), [
        ("Teslim Adresi/yeri:", True),
        (" {{TESLIM_YERI}} veya Kiracının sözleşme süresince Kiralayana elektronik posta ile bildireceği "
         "diğer teslim/çalışma adresidir.", False)])
    paragraf_yaz(bul(doc, "Teslim Tarihi:"), [
        ("Teslim Tarihi:", True),
        (" Kiralayan tarafından Makinanın sağlam ve çalışır vaziyette Kiracıya teslim edileceği, "
         "en geç {{BASLANGIC}} olan tarihi ifade eder. Kiracının sözleşme süresince elektronik posta ile "
         "sevk talebinde bulunması halinde Makina, talebi takip eden en geç {{SEVK_SURESI}} içinde Kiracının "
         "belirttiği adrese teslim edilir; bu süre mutlak olup aşılması Madde 7.3 kapsamında gecikme sayılır. "
         "Fiili teslim tarihi Teslim Tutanağı ile belirlenir.", False)])

    # --- MADDE 3: MAKİNA BİLGİLERİ -----------------------------------------
    p32 = bul(doc, "3.2.")
    paragraf_yaz(p32, [
        ("3.2. ", True),
        ("Kiralamaya konu iş makinalarının bilgileri ve makina başına aylık kira bedelleri aşağıdaki "
         "tabloda gösterilmiştir. Sözleşmede geçen \"Makina\" ifadesi tabloda yer alan makinaların her "
         "birini ayrı ayrı ve tamamını birlikte kapsar; sözleşme hükümleri (kira kesintisi, gecikme cezası, "
         "ikame ve fesih hakları dahil) her makina için ayrı ayrı uygulanır. Tabloda belirtilen teknik "
         "özellikler Kiralayanın taahhüdü niteliğinde olup ataşmanlar makinanın ayrılmaz parçası ve kira "
         "bedeline dahildir.", False)])
    for etiket in ("Plakası", "Cinsi", "Markası", "Modeli", "Model Yılı", "Şasi No"):
        paragraf_sil(bul(doc, etiket))
    paragraf_ekle_sonra(p32, [("{{ARAC_TABLOSU}}", False)])

    # --- MADDE 4: SÜRE -----------------------------------------------------
    paragraf_yaz(bul(doc, "4.1."), [
        ("4.1. ", True),
        ("Sözleşme imza tarihinde yürürlüğe girer. Kira süresi, Teslim Tutanağında belirtilen Kira "
         "Ödemelerinin Başlangıç Tarihinde başlar ve ", False), ("{{BITIS}}", True),
        (" tarihinde hiçbir ihbara gerek kalmaksızın sona erer.", False)])
    paragraf_ekle_sonra(bul(doc, "4.3."), [
        ("4.3.1. ", True),
        ("Kiracı erken fesih ve iade iradesini, iade tarihinden en az 3 (üç) gün önce Madde 1.4'te "
         "belirtilen elektronik posta adresine bildirir. Erken fesih hakkı makinaların tamamı veya bir "
         "kısmı için kullanılabilir. Kira bedeli, iade tutanağının düzenlendiği tarihe kadar Madde 5.1.2 "
         "uyarınca hesaplanır. İadeye ilişkin nakliye ve sair masraflar Madde 5.3 uyarınca Kiralayana aittir.", False)])

    # --- MADDE 5: ÜCRET ----------------------------------------------------
    p51 = paragraf_yaz(bul(doc, "5.1."), [
        ("5.1. ", True),
        ("Makina başına aylık kira bedelleri Madde 3.2'deki tabloda gösterilmiş olup toplam aylık kira "
         "bedeli net ", False), ("{{ARA_TOPLAM}} ({{ARA_TOPLAM_YAZI}})", True),
        (" + %{{KDV_ORANI}} KDV ", False), ("{{KDV}}", True), (" olmak üzere ", False),
        ("KDV dahil {{GENEL_TOPLAM}} ({{GENEL_TOPLAM_YAZI}})", True),
        (" olarak belirlenmiştir.", False)])
    paragraf_ekle_sonra(p51, [
        ("5.1.2. ", True),
        ("Ay içinde kısmi kullanımlarda kira bedeli takvim günü esasıyla (aylık kira bedeli / 30 x "
         "kullanılan gün) hesaplanır. Bu sözleşmede geçen \"günlük kira bedeli\" ifadesi, kira kesintisi "
         "ve gecikme cezası hesaplarında ilgili makinanın aylık kira bedelinin 26 (yirmi altı)'ya "
         "bölünmesiyle bulunan tutardır.", False)])
    paragraf_ekle_sonra(p51, [
        ("5.1.1. ", True),
        ("Aylık kira bedeli; karla mücadele ve kış bakım çalışmalarının niteliği gereği gece, hafta sonu, "
         "resmi tatil ve dini bayramlar dahil haftanın 7 günü 24 saat hazır bekleme ve çalışmayı kapsar. "
         "Çalışma saati, mesai veya çalışılan gün sayısı gerekçe gösterilerek fazla çalışma bedeli veya "
         "ilave ücret talep edilemez.", False)])

    p52 = bul(doc, "5.2.")
    metin_degistir(p52, "işbu süreler toplam kira süresine otomatikman eklenecektir.",
                   "Kiracı dilerse, yazılı bildirimi ile işbu süreleri kira süresine ekletebilir; "
                   "böyle bir bildirim yapılmadıkça kira süresi uzamaz.")
    p52.add_run(" Arıza, bakım ve ikame makinaya ilişkin süreler, kira kesintisi ve gecikme cezası Madde 7'de "
                "düzenlenmiş olup Madde 5.2, 6.1.4 ve 7 hükümleri birbirini tamamlar nitelikte birlikte "
                "uygulanır. Aynı süre için kira kesintisi, gecikme cezası ve (Kiracının talebi halinde) süre "
                "uzatımı birbirini engellemeksizin birlikte uygulanır.")
    paragraf_ekle_sonra(p52, [
        ("5.2.1. ", True),
        ("Kira bedelleri yalnızca Kiralayan adına açılmış ", False), ("{{K_BANKA}}", True),
        (" nezdindeki ", False), ("{{K_IBAN}}", True),
        (" IBAN numaralı hesaba ödenir. Kiralayan, hesap bilgisindeki değişikliği imzalı yazı ile bildirmedikçe "
         "bu hesaba yapılan ödemeler geçerli ve borçtan kurtarıcıdır; üçüncü kişi adına açılmış hesaplara "
         "ödeme talep edilemez.", False)])
    bul(doc, "5.4.").add_run(
        " Kiracı mahsup ettiği tutarı dayanak belgeleriyle birlikte Kiralayanın elektronik posta adresine "
        "bildirir; bu bildirim mahsubun geçerlilik şartı olmayıp yalnızca bilgilendirme amacı taşır.")

    # --- MADDE 6 -----------------------------------------------------------
    paragraf_ekle_sonra(bul(doc, "6.1.4."), [
        ("6.1.5. ", True),
        ("Karla mücadele çalışmalarında kullanılan tuz, solüsyon ve sair kimyasalların makinada neden "
         "olduğu korozyon ile kar bıçağı, bıçak lastiği/kesici ağız, tuz serici ekipmanı ve lastiklerde "
         "meydana gelen aşınma, makinanın sözleşmeye uygun olağan kullanımından kaynaklanan aşınma "
         "niteliğindedir. Bunlardan dolayı Kiracı sorumlu tutulamaz; Kiralayan tarafından sağlanan "
         "makina ve ataşmanlara ilişkin bakım, onarım ve değişim giderleri Kiralayana aittir.", False)])

    # --- MADDE 7 -----------------------------------------------------------
    metin_degistir(bul(doc, "7.1."), "23.02.2026", "{{BASLANGIC}}")
    metin_degistir(bul(doc, "- Araç üzerinde 3."), "mezkur mübayenetin", "mezkur durumun")

    # --- MADDE 8: İMM terimi düzeltmesi -----------------------------------
    metin_degistir(bul(doc, "8.1."), "İşveren Mali Mesuliyet (İMM)", "İhtiyari Mali Mesuliyet (İMM)")

    # --- MADDE 10, 13, 14 --------------------------------------------------
    metin_degistir(bul(doc, "İşbu madde sahilinde"), "madde sahilinde", "madde dahilinde")
    metin_degistir(bul(doc, "Teslim ve İade tutanakları"), "Teslim ve İade tutanakları",
                   "Makina Teslim ve İade Tutanakları (Ek-1)")
    sigorta = bul(doc, "Sigorta poliçeleri")
    paragraf_ekle_sonra(sigorta, [("Kiralayan şahıs ise nüfus cüzdanı/kimlik kartı fotokopisi", False)])
    paragraf_ekle_sonra(sigorta, [("Kiralayana ait banka hesap (IBAN) teyit belgesi", False)])
    paragraf_ekle_sonra(sigorta, [("Kiralayanın vergi levhası (fatura düzenleyen kiralayanlar için)", False)])
    paragraf_yaz(bul(doc, "İşbu 14 madde"), [
        ("İşbu 14 (on dört) madde ve eklerinden müteşekkil İş Makinası Kiralama Sözleşmesi imza tarihi "
         "itibarıyla yürürlüğe girmekte olup ", False), ("{{IMZA_TARIHI}}", True),
        (" tarihinde taraflar arasında 2 (iki) nüsha olarak tanzim edilmiş ve imza altına alınmıştır.", False)])

    # --- İmza tablosu ------------------------------------------------------
    imza = doc.tables[0]
    hucre = imza.rows[1].cells
    paragraf_yaz(hucre[0].paragraphs[0], [("{{K_AD}}", True)])
    paragraf_yaz(hucre[1].paragraphs[0], [("{{R_UNVAN}}", True)])
    for c, alan in ((imza.rows[2].cells[0], "{{K_IMZACI}}"), (imza.rows[2].cells[1], "{{R_IMZACI}}")):
        paragraf_yaz(c.paragraphs[0], [(alan, False)])
        c.add_paragraph("İmza / Kaşe:")
    # Madde 14 ve imza tablosu sayfa arasında bölünmesin
    bul(doc, "MADDE-14").paragraph_format.keep_with_next = True
    bul(doc, "İşbu 14").paragraph_format.keep_with_next = True
    for i, satir in enumerate(imza.rows):
        trPr = satir._tr.get_or_add_trPr()
        trPr.append(trPr.makeelement(qn("w:cantSplit"), {}))
        for c in satir.cells:
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_with_next = i < len(imza.rows) - 1

    # İmza tablosundan sonraki boş paragraflar tutanaklardan önce boş sayfa yaratıyordu
    sonraki = imza._tbl.getnext()
    while sonraki is not None and sonraki.tag == qn("w:p"):
        sil, sonraki = sonraki, sonraki.getnext()
        sil.getparent().remove(sil)
    # Ek-1 tutanakları buraya üretilecek
    doc.add_paragraph("{{TUTANAKLAR}}")

    # Hukukçu yorumlarını ve renkli vurguları temiz çıktı için kaldır (orijinal dosyada duruyor)
    govde = doc.element.body
    for etiket in ("w:commentRangeStart", "w:commentRangeEnd", "w:highlight"):
        for e in list(govde.iter(qn(etiket))):
            e.getparent().remove(e)
    for ref in list(govde.iter(qn("w:commentReference"))):
        run = ref.getparent()
        run.getparent().remove(run)
    for rel in list(doc.part.rels.values()):
        if rel.reltype.endswith(("/comments", "/commentsExtended", "/commentsIds", "/commentsExtensible")):
            doc.part.drop_rel(rel.rId)

    kurumsal_bicim(doc)
    doc.save(HEDEF)
    print(f"Şablon oluşturuldu: {HEDEF}")


if __name__ == "__main__":
    main()
