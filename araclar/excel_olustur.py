"""Boş ARACLAR.xlsx veri giriş dosyasını oluşturur. Var olan dosyanın üzerine YAZMAZ."""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

KOK = Path(__file__).resolve().parent.parent
HEDEF = Path(sys.argv[1]) if len(sys.argv) > 1 else KOK / "ARACLAR.xlsx"

BASLIK = Font(bold=True, color="FFFFFF")
DOLGU = PatternFill("solid", fgColor="1F4E78")
ZORUNLU = PatternFill("solid", fgColor="C00000")

AYARLAR = [
    ("Kiracı Unvanı", "ÖZTAÇ PETROL A.Ş-Hİ-KA NAK.LTD.ŞTİ -İŞ ORTAKLIĞI"),
    ("Kiracı Adresi", "D 100 GAP MAH. TURGUT ÖZAL BLV. GOLDEN PLAZA 52 521 Batman Merkez/Batman"),
    ("Kiracı Telefon", ""),
    ("Kiracı E-posta", ""),
    ("Kiracı Vergi Dairesi", "BATMAN VERGİ DAİRESİ MÜD."),
    ("Kiracı VKN", "7060409340"),
    ("Kiracı İmza Yetkilisi (Ad Soyad / Unvan)", ""),
    ("Sevk Talebinden Sonra Teslim Süresi", "48 (kırk sekiz) saat"),
    ("KDV Oranı (%)", 20),
]

# (başlık, genişlik, zorunlu mu, sayı biçimi)
SOZLESME_SUTUN = [
    ("Sözleşme No", 12, True, "@"),
    ("Üret (E/H)", 9, True, None),
    ("Kiralayan Adı Soyadı / Unvanı", 34, True, None),
    ("TCKN (şahıs) / VKN (şirket)", 16, True, "@"),
    ("Vergi Dairesi", 18, False, None),
    ("Kiralayan Adresi", 40, True, None),
    ("Telefon", 15, False, "@"),
    ("E-posta", 24, False, None),
    ("Şirket Yetkilisi (şahıssa boş)", 24, False, None),
    ("Banka Adı / Şube", 22, True, None),
    ("IBAN", 32, True, "@"),
    ("Teslim (Başlangıç) Tarihi", 14, True, "DD.MM.YYYY"),
    ("Bitiş Tarihi", 14, True, "DD.MM.YYYY"),
    ("Teslim Yeri / Adresi", 36, True, None),
    ("İmza Tarihi (boşsa elle)", 14, False, "DD.MM.YYYY"),
]
ARAC_SUTUN = [
    ("Sözleşme No", 12, True, "@"),
    ("Plaka", 12, False, "@"),
    ("Cinsi", 18, True, None),
    ("Marka", 14, True, None),
    ("Model", 14, True, None),
    ("Model Yılı", 10, False, "0"),
    ("Şasi No", 22, True, "@"),
    ("Teknik Özellikler (kapasite, motor gücü vb.)", 30, False, None),
    ("Ataşmanlar (virgülle ayırın)", 34, False, None),
    ("Aylık Kira (KDV hariç, TL)", 18, True, '#,##0.00'),
]

YARDIM = """NASIL KULLANILIR
1) 'Ayarlar' sayfasında şirket (Kiracı) bilgilerini bir kez kontrol edin / tamamlayın.
2) 'Sözleşmeler' sayfasına her kiralayan firma/kişi için BİR satır ekleyin. Sözleşme No benzersiz olmalı (örn. 2026-001).
3) 'Araçlar' sayfasına o kiralayandan alınan HER araç için bir satır ekleyin; Sözleşme No ile bağlayın.
   Bir sözleşmede 1 araç da olabilir, 10 araç da.
4) Üretmek istediğiniz sözleşmelerde 'Üret (E/H)' = E yapın, dosyayı KAYDEDİN.
5) Klasördeki SOZLESME_URET.bat dosyasına çift tıklayın. Çıktılar CIKTI klasörüne Word + PDF olarak düşer.
Kırmızı başlıklı sütunlar zorunludur; eksik varsa o sözleşme üretilmez ve ekranda nedeni yazılır.
Sözleşme metnini değiştirmek için SABLON.docx dosyasını Word ile düzenleyin; {{...}} alanlarına dokunmayın."""


def sayfa_kur(ws, sutunlar):
    for i, (ad, gen, zorunlu, bicim) in enumerate(sutunlar, 1):
        h = ws.cell(row=1, column=i, value=ad)
        h.font, h.fill = BASLIK, ZORUNLU if zorunlu else DOLGU
        h.alignment = Alignment(wrap_text=True, vertical="center")
        harf = h.column_letter
        ws.column_dimensions[harf].width = gen
        if bicim:
            for r in range(2, 501):
                ws[f"{harf}{r}"].number_format = bicim
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def main():
    if HEDEF.exists():
        sys.exit(f"{HEDEF} zaten var, üzerine yazılmadı.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Nasıl Kullanılır"
    for i, satir in enumerate(YARDIM.splitlines(), 1):
        ws.cell(row=i, column=1, value=satir).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Ayarlar")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 75
    for i, (k, v) in enumerate(AYARLAR, 1):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v).number_format = "@" if k != "KDV Oranı (%)" else "0"

    ws = wb.create_sheet("Sözleşmeler")
    sayfa_kur(ws, SOZLESME_SUTUN)
    dv = DataValidation(type="list", formula1='"E,H"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B2:B500")

    sayfa_kur(wb.create_sheet("Araçlar"), ARAC_SUTUN)
    wb.active = 2
    wb.save(HEDEF)
    print(f"Oluşturuldu: {HEDEF}")


if __name__ == "__main__":
    main()
