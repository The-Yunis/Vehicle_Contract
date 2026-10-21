"""SABLON.docx üzerinde metin değişiklikleri (v1.6.0+).

sablon_olustur.py şablonu şirketin ilk taslağından üretmişti; o taslağın yerini sonradan şirketin
doldurduğu örnek sözleşme aldı. Bu yüzden yeni metin değişiklikleri hazır SABLON.docx'e uygulanır.
Her adım tekrar çalıştırılabilir (aynı değişikliği iki kez eklemez).

Kullanım:  python araclar/sablon_duzenle.py
"""
import re
import sys
from pathlib import Path

import docx

from bicim import altbilgi
from ortak import araliga_yaz, metin_degistir, paragraf_ekle_sonra, paragraf_yaz

KOK = Path(__file__).resolve().parent.parent
SABLON = KOK / "SABLON.docx"

# Kar küreme kamyonu iş makinası değil araç olduğundan şirket metni "araç ve iş makinası" diye düzeltti.
IS_MAKINASI = re.compile(r"(?<![Aa]raç ve )\b(İş|iş) ([Mm]akin\w*)")


def tum_paragraflar(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for satir in t.rows:
            for c in satir.cells:
                yield from c.paragraphs


def arac_ve_is_makinasi(doc):
    sayi = 0
    for p in tum_paragraflar(doc):
        # belgenin adı ("İş Makinası Kiralama Sözleşmesi") olduğu gibi kalır
        eslesmeler = [m for m in IS_MAKINASI.finditer(p.text)
                      if not p.text[m.end():].lstrip().startswith("Kiralama Sözleşmesi")]
        for m in reversed(eslesmeler):   # sondan başa: önceki konumlar kaymasın
            ek = "Araç ve " if m.group(1) == "İş" else "araç ve "
            araliga_yaz(p, m.start(1), m.end(1), ek + m.group(1))
            sayi += 1
    return sayi


def ve_tekrari(doc):
    """'... Tutanağı" ve araç ve iş makinasının' → '... Tutanağı" ile araç ve iş makinasının'."""
    for p in doc.paragraphs:
        metin_degistir(p, '" ve araç ve iş makinasının', '" ile araç ve iş makinasının')


def odeme_alanlari(doc):
    """5.2 ödeme şeklini ve 5.2.1 ödeme hükmünü sözleşmeye göre doldurulan alanlara çevirir."""
    degisti = False
    for p in doc.paragraphs:
        if p.text.startswith("5.2.\t") and "banka yoluyla ödenecektir" in p.text:
            metin_degistir(p, "banka yoluyla ödenecektir", "{{ODEME_SEKLI}} ödenecektir")
            degisti = True
        if p.text.startswith("5.2.1.\t") and "{{ODEME_EK}}" not in p.text:
            paragraf_yaz(p, [("5.2.1.\t", True), ("{{ODEME_EK}}", False)])
            degisti = True
    return degisti


def kiralama_turu(doc):
    """3.1'deki 'operatörsüz (şoförsüz)' ve 6.2.6 personel hükmü sözleşmeye göre doldurulur."""
    for p in doc.paragraphs:
        if p.text.startswith("3.1.\t") and "operatörsüz (şoförsüz)" in p.text:
            metin_degistir(p, "operatörsüz (şoförsüz)", "{{KIRALAMA_TURU}}")
        if p.text.startswith("6.2.6.\t") and "{{OPERATOR_HUKMU}}" not in p.text:
            paragraf_yaz(p, [("6.2.6.\t", True), ("{{OPERATOR_HUKMU}}", False)])


def kimlik_satirlari(doc):
    """Kiralayan: Vergi No her zaman yazılır; T.C. Kimlik No isteğe bağlı ayrı satır (boşsa satır silinir)."""
    for p in doc.paragraphs:
        m = re.search(r"\s*\{\{K_KIMLIK\}\}", p.text)
        if m:
            araliga_yaz(p, m.start(), m.end(), "")
    if any("{{?K_TCKN}}" in p.text for p in doc.paragraphs):
        return
    vn = next(p for p in doc.paragraphs if "{{K_VN}}" in p.text)
    paragraf_ekle_sonra(vn, [("T.C. Kimlik No", True), ("\t: ", True), ("{{?K_TCKN}}", False)])


def main():
    doc = docx.Document(SABLON)
    sayi = arac_ve_is_makinasi(doc)
    ve_tekrari(doc)
    odeme = odeme_alanlari(doc)
    kiralama_turu(doc)
    kimlik_satirlari(doc)
    altbilgi(doc)   # yalnızca sayfa numarası
    doc.save(SABLON)
    print(f'"araç ve iş makinası": {sayi} yer | ödeme alanları: {"eklendi" if odeme else "zaten var"}')
    kalan = [p.text[:70] for p in tum_paragraflar(doc) if re.search(r"(?<!ve )\b[İi]ş [Mm]akin", p.text)]
    if kalan:
        print("Dokunulmayanlar (belge adı vb.):", *kalan, sep="\n  ")


if __name__ == "__main__":
    sys.exit(main())
